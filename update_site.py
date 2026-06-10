#!/usr/bin/env python3
"""
update_site.py — Rebuild index.html from any resume file and push to GitHub.

Usage:
    python3 update_site.py <path-to-resume>

Supported formats: .pptx  .pdf  .docx  .txt

Examples:
    python3 update_site.py ~/Downloads/HangNguyen_DataScience.pptx
    python3 update_site.py ~/Downloads/resume.pdf
    python3 update_site.py ~/Downloads/cv.docx
    python3 update_site.py ~/Downloads/bio.txt
"""

import subprocess, sys, os, re
from pathlib import Path

SITE_DIR    = Path.home() / "HangNguyen_Website"
HTML_PATH   = SITE_DIR / "index.html"
PROFILE_IMG = "profile.jpg"
GITHUB_URL  = "https://hangnguyen4142.github.io/hang-nguyen-site/"

# ── 1. ARG CHECK ──────────────────────────────────────────────────────────────
if len(sys.argv) < 2:
    print("Usage: python3 update_site.py <path-to-resume>")
    print("Example: python3 update_site.py ~/Downloads/MyResume.pptx")
    sys.exit(1)

input_path = Path(sys.argv[1]).expanduser().resolve()
if not input_path.exists():
    print(f"❌  File not found: {input_path}")
    sys.exit(1)

ext = input_path.suffix.lower()
if ext not in {".pptx", ".pdf", ".docx", ".txt"}:
    print(f"❌  Unsupported format '{ext}'. Supported: .pptx .pdf .docx .txt")
    sys.exit(1)

print(f"📄  Reading {input_path.name}...")

# ── 2. AUTO-INSTALL DEPS ──────────────────────────────────────────────────────
def pip_install(pkg):
    print(f"    Installing {pkg}...")
    subprocess.run([sys.executable, "-m", "pip", "install", pkg,
                    "--break-system-packages", "-q"], check=True)

# ── 3. EXTRACT TEXT ───────────────────────────────────────────────────────────
shapes = {}   # for .pptx structured parsing
raw_text = ""

if ext == ".pptx":
    try:
        from pptx import Presentation
    except ImportError:
        pip_install("python-pptx"); from pptx import Presentation
    try:
        prs = Presentation(str(input_path))
    except Exception as e:
        print(f"❌  Cannot read PPTX: {e}")
        print("    Remove encryption first: PowerPoint → File → Info → Protect → Remove Protection")
        sys.exit(1)
    for slide in prs.slides:
        for shape in slide.shapes:
            if shape.has_text_frame:
                text = "\n".join(p.text.strip() for p in shape.text_frame.paragraphs if p.text.strip())
                if text:
                    shapes[shape.name] = text
                    raw_text += text + "\n"

elif ext == ".pdf":
    try:
        import pdfplumber
    except ImportError:
        pip_install("pdfplumber"); import pdfplumber
    with pdfplumber.open(str(input_path)) as pdf:
        for page in pdf.pages:
            t = page.extract_text()
            if t:
                raw_text += t + "\n"

elif ext == ".docx":
    try:
        from docx import Document
    except ImportError:
        pip_install("python-docx"); from docx import Document
    doc = Document(str(input_path))
    for para in doc.paragraphs:
        if para.text.strip():
            raw_text += para.text.strip() + "\n"

elif ext == ".txt":
    raw_text = input_path.read_text(encoding="utf-8", errors="ignore")

if not raw_text.strip():
    print("❌  No text could be extracted from this file.")
    sys.exit(1)

print(f"    ✓ Extracted {len(raw_text):,} characters")

# ── 4. PARSE CONTENT ──────────────────────────────────────────────────────────
name = ""; job_title = ""; location = ""; background = ""
edu = ""; certs = []; skills = {}; jobs = []

if ext == ".pptx" and shapes:
    # Structured parsing using PowerPoint shape names
    name_block = shapes.get("Text Box 7", "")
    job_title  = shapes.get("Title 1", "")
    exp_block  = shapes.get("Rectangle 9", "")
    job_block  = shapes.get("Rectangle 13", "")

    name_lines = name_block.splitlines()
    name     = name_lines[0] if name_lines else ""
    location = name_lines[2] if len(name_lines) > 2 else ""

    cur = None
    for line in exp_block.splitlines():
        line = line.strip()
        if line.startswith("EXPERIENCE:"):
            background = line[len("EXPERIENCE:"):].strip()
        elif "EDUCATION" in line and "CERTIF" in line:
            cur = "edu"
        elif "TECHNICAL" in line and "EXPERIENCE" in line:
            cur = "tech"
        elif cur == "edu" and line:
            if any(x in line for x in ["B.S", "M.S", "Bachelor", "Master"]):
                edu = line
            else:
                certs.append(line)
        elif cur == "tech" and line and ":" in line:
            k, _, v = line.partition(":")
            skills[k.strip()] = v.strip()

    # Parse job entries
    for section in re.split(r'\n(?=[A-Z][^\n•]{5,}(?:\n|$))', job_block.strip()):
        lines = [l.strip() for l in section.strip().splitlines() if l.strip()]
        if not lines:
            continue
        header = lines[0]
        sep = " - " if " - " in header else (" \u2013 " if " \u2013 " in header else None)
        company, role = ("", header) if not sep else (header.split(sep,1)[0].strip(), header.split(sep,1)[1].strip())
        bullets = [l.lstrip("\u2022").strip() for l in lines[1:] if l.strip()]
        if bullets:
            jobs.append({"role": role, "company": company, "bullets": bullets})

else:
    # Generic heuristic parsing for .pdf / .docx / .txt
    lines = [l.strip() for l in raw_text.splitlines() if l.strip()]
    name = lines[0] if lines else "Name"
    cur = None
    cur_job = None

    for line in lines[1:]:
        low = line.lower()
        if re.match(r'^(experience|employment|work history)\s*$', low):
            cur = "exp"; continue
        elif re.match(r'^(education|academic)\s*$', low):
            cur = "edu"; continue
        elif re.match(r'^(certif|license)\w*\s*$', low):
            cur = "cert"; continue
        elif re.match(r'^(skill|technical|competenc)\w*\s*$', low):
            cur = "skill"; continue
        elif re.match(r'^(summary|profile|background|objective)\s*$', low):
            cur = "bg"; continue

        if cur == "bg":
            background += line + " "
        elif cur == "edu" and not edu:
            edu = line
        elif cur == "cert":
            certs.append(line)
        elif cur == "skill" and ":" in line:
            k, _, v = line.partition(":"); skills[k.strip()] = v.strip()
        elif cur == "exp":
            if not line.startswith(("•", "-")) and len(line) < 100:
                if cur_job and cur_job["bullets"]:
                    jobs.append(cur_job)
                cur_job = {"role": line, "company": "", "bullets": []}
            elif cur_job:
                cur_job["bullets"].append(line.lstrip("•-").strip())

    if cur_job and cur_job.get("bullets"):
        jobs.append(cur_job)
    background = background.strip()

# ── 5. HTML HELPERS ───────────────────────────────────────────────────────────
BG_SVG = (
    "url(\"data:image/svg+xml,%3Csvg xmlns='http://www.w3.org/2000/svg' width='160' height='120' "
    "viewBox='0 0 160 120'%3E%3Cg opacity='0.13' fill='%231a8c8c'%3E%3Cellipse cx='62' cy='52' rx='28' ry='18'/%3E"
    "%3Cellipse cx='42' cy='58' rx='18' ry='14'/%3E%3Cellipse cx='82' cy='58' rx='18' ry='14'/%3E"
    "%3Crect x='24' y='58' width='76' height='16' rx='4'/%3E%3Crect x='52' y='55' width='20' height='16' rx='3'/%3E"
    "%3Cpath d='M56 55 L56 50 C56 44 68 44 68 50 L68 55' fill='none' stroke='%231a8c8c' stroke-width='3.5' stroke-linecap='round'/%3E"
    "%3Ccircle cx='62' cy='61' r='2.5' fill='%23e4eff4'/%3E%3Crect x='61' y='63' width='2' height='4' rx='1' fill='%23e4eff4'/%3E"
    "%3C/g%3E%3C/svg%3E\")"
)

def esc(s):
    return str(s).replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

def cert_html(items):
    out = ""
    for c in items:
        cls = "cert-badge in-progress" if "in progress" in c.lower() else "cert-badge"
        out += f'      <span class="{cls}">{esc(c)}</span>\n'
    return out

def skill_html(text):
    tags = [t.strip().lstrip("•").strip() for t in text.replace("•","\n").splitlines() if t.strip().lstrip("•").strip()]
    return "".join(f'<span class="skill-tag">{esc(t)}</span>' for t in tags if t)

def jobs_html(job_list):
    out = ""
    for j in job_list:
        out += (
            '      <div class="job">\n'
            '        <div class="job-header">\n'
            f'          <h3>{esc(j["role"])}</h3>\n'
            f'          <span class="company">{esc(j["company"])}</span>\n'
            '        </div>\n'
            '        <ul>\n'
        )
        for b in j["bullets"]:
            out += f"          <li>{esc(b)}</li>\n"
        out += "        </ul>\n      </div>\n\n"
    return out

skills_sidebar = "".join(
    f'<p class="skill-cat">{esc(k)}</p><div class="skills-grid">{skill_html(v)}</div>'
    for k, v in skills.items()
)

tagline_parts = [p for p in [location, background[:120] + ("..." if len(background) > 120 else "")] if p]
tagline = "  ·  ".join(tagline_parts)

# ── 6. BUILD HTML ─────────────────────────────────────────────────────────────
print("🔨  Building index.html...")

CSS = """
    :root {
      --navy: #0d1b2a; --teal: #1a8c8c; --teal-light: #23b5b5;
      --accent: #e8f4f4; --white: #ffffff; --text: #1e2d3d;
      --muted: #5a7080; --border: #d0e4e4;
    }
    * { box-sizing: border-box; margin: 0; padding: 0; }
    body {
      font-family: 'Segoe UI', system-ui, sans-serif;
      color: var(--text); line-height: 1.6;
      background-color: #e4eff4;
      background-image: BG_SVG_PLACEHOLDER;
      background-size: 240px 180px;
    }
    header { background: var(--navy); color: #fff; padding: 2.5rem 3rem; display: flex; align-items: center; gap: 2rem; }
    header img { width: 110px; height: 110px; border-radius: 50%; object-fit: cover; border: 3px solid var(--teal-light); flex-shrink: 0; }
    header .intro h1 { font-size: 2rem; font-weight: 700; }
    header .intro p.title { font-size: 1.05rem; color: var(--teal-light); margin-top: .3rem; font-weight: 500; }
    header .intro p.tagline { margin-top: .6rem; font-size: .9rem; color: #a8c8d0; max-width: 620px; }
    .page { display: grid; grid-template-columns: 280px 1fr; max-width: 1100px; margin: 2rem auto; gap: 1.5rem; padding: 0 1.5rem 2rem; }
    aside { display: flex; flex-direction: column; gap: 1.5rem; }
    .card { background: var(--white); border-radius: 10px; padding: 1.4rem 1.3rem; box-shadow: 0 2px 8px rgba(0,0,0,.07); }
    .card h2 { font-size: .78rem; text-transform: uppercase; letter-spacing: 1.2px; color: var(--teal); border-bottom: 2px solid var(--border); padding-bottom: .5rem; margin-bottom: .9rem; }
    .card p { font-size: .88rem; color: var(--muted); line-height: 1.65; }
    .card ul { list-style: none; font-size: .88rem; color: var(--text); }
    .card ul li { padding: .3rem 0; border-bottom: 1px solid #eef4f4; display: flex; gap: .5rem; }
    .card ul li:last-child { border-bottom: none; }
    .card ul li::before { content: '▸'; color: var(--teal); flex-shrink: 0; margin-top: 1px; }
    .cert-badge { display: inline-block; background: var(--accent); color: var(--teal); border: 1px solid var(--border); border-radius: 20px; padding: .3rem .75rem; font-size: .78rem; margin: .25rem .2rem 0 0; font-weight: 500; }
    .cert-badge.in-progress { background: #fff8e6; color: #9c7400; border-color: #f0d070; }
    .skills-grid { display: flex; flex-wrap: wrap; gap: .4rem; margin-top: .4rem; }
    .skill-tag { background: #edf7f7; color: var(--teal); border-radius: 4px; padding: .2rem .55rem; font-size: .77rem; font-weight: 500; border: 1px solid var(--border); }
    .skill-cat { font-size: .8rem; color: var(--teal); font-weight: 600; margin: .7rem 0 .4rem; }
    main { display: flex; flex-direction: column; gap: 1.5rem; }
    .section { background: var(--white); border-radius: 10px; padding: 1.6rem 1.8rem; box-shadow: 0 2px 8px rgba(0,0,0,.07); }
    .section h2 { font-size: .78rem; text-transform: uppercase; letter-spacing: 1.2px; color: var(--teal); border-bottom: 2px solid var(--border); padding-bottom: .5rem; margin-bottom: 1.2rem; }
    .job { margin-bottom: 1.8rem; }
    .job:last-child { margin-bottom: 0; }
    .job-header { display: flex; justify-content: space-between; align-items: flex-start; flex-wrap: wrap; gap: .3rem; }
    .job-header h3 { font-size: 1rem; font-weight: 700; color: var(--navy); }
    .job-header .company { font-size: .85rem; color: var(--teal); font-weight: 600; }
    .job ul { padding-left: 1.2rem; font-size: .88rem; color: #2e4050; line-height: 1.7; }
    .job ul li { margin-bottom: .5rem; }
    footer { text-align: center; padding: 1.5rem; font-size: .8rem; color: var(--muted); }
    @media (max-width: 750px) {
      header { flex-direction: column; text-align: center; padding: 2rem 1.5rem; }
      .page { grid-template-columns: 1fr; }
    }
""".replace("BG_SVG_PLACEHOLDER", BG_SVG)

edu_card    = f'<div class="card"><h2>Education</h2><ul><li>{esc(edu)}</li></ul></div>' if edu else ""
cert_card   = f'<div class="card"><h2>Certifications</h2>\n{cert_html(certs)}</div>' if certs else ""
skills_card = f'<div class="card"><h2>Technical Skills</h2>{skills_sidebar}</div>' if skills else ""
footer_text = esc(name) + (" &middot; " + esc(job_title) if job_title else "")

html = f"""<!DOCTYPE html>
<html lang="en">
<head>
  <meta charset="UTF-8" />
  <meta name="viewport" content="width=device-width, initial-scale=1.0" />
  <title>{esc(name)} | {esc(job_title)}</title>
  <style>{CSS}  </style>
</head>
<body>
<header>
  <img src="{PROFILE_IMG}" alt="{esc(name)}" />
  <div class="intro">
    <h1>{esc(name)}</h1>
    <p class="title">{esc(job_title)}</p>
    <p class="tagline">{esc(tagline)}</p>
  </div>
</header>
<div class="page">
  <aside>
    <div class="card">
      <h2>Professional Background</h2>
      <p>{esc(background)}</p>
    </div>
    {edu_card}
    {cert_card}
    {skills_card}
  </aside>
  <main>
    <div class="section">
      <h2>Selected Experience</h2>
{jobs_html(jobs)}    </div>
  </main>
</div>
<footer>&copy; 2026 {footer_text}</footer>
</body>
</html>
"""

HTML_PATH.write_text(html, encoding="utf-8")
print("✅  index.html updated")

# ── 7. PUSH TO GITHUB ─────────────────────────────────────────────────────────
print("🚀  Pushing to GitHub...")
os.chdir(SITE_DIR)
subprocess.run(["git", "add", "index.html"], check=True)
result = subprocess.run(["git", "diff", "--cached", "--quiet"])
if result.returncode == 0:
    print("✅  No changes — site is already up to date.")
else:
    subprocess.run(["git", "commit", "-m", f"Update site from {input_path.name}"], check=True)
    subprocess.run(["git", "push"], check=True)
    print(f"✅  Done! Live at: {GITHUB_URL}")
    print("    (GitHub Pages rebuilds in ~1 minute)")
